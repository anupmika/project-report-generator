"""Format parser for custom project report templates."""

from typing import Any

from docx import Document


class ReportFormatParser:
    """Parse a DOCX project report format template."""

    def __init__(self, template_path: str):
        self.template_path = template_path
        self.doc = Document(template_path)
        self.sections: list[dict[str, Any]] = []

    def parse(self) -> dict[str, Any]:
        """Parse the template to extract structure."""
        structure: dict[str, Any] = {
            'sections': [],
            'headings': [],
            'placeholders': [],
        }

        for para in self.doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            # Check if it's a heading
            style_name = para.style.name.lower()
            if style_name.startswith('heading'):
                structure['headings'].append({
                    'level': style_name,
                    'text': text
                })
                structure['sections'].append({
                    'type': 'heading',
                    'level': style_name,
                    'content': text
                })
            elif text.startswith(('{{', '[')):
                structure['placeholders'].append(text)
                structure['sections'].append({
                    'type': 'placeholder',
                    'content': text
                })
            else:
                structure['sections'].append({
                    'type': 'paragraph',
                    'content': text
                })

        # Also extract tables
        for table in self.doc.tables:
            structure['sections'].append({
                'type': 'table',
                'rows': len(table.rows),
                'cols': len(table.columns) if table.rows else 0
            })

        self.sections = structure['sections']
        return structure

    def get_structure(self) -> list[dict[str, Any]]:
        """Get parsed structure."""
        if not self.sections:
            self.parse()
        return self.sections
